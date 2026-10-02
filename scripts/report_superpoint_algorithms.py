#!/usr/bin/env python3
"""Consolidate the SPT/EZ-SP partition-kernel experiments into JSON and Excel."""
from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook

PROJECT = Path(__file__).resolve().parents[1]
ROOT = PROJECT / "outputs/dualcrown3d_superpoint_v1"
RUNS = (
    ("frozen_features_plus_prior_affinity", "stage2_spt_ezsp_partition_pilot_v2"),
    ("instance_embedding_unit_weights", "stage2_instance_partition_unit"),
    ("instance_embedding_affinity_weights", "stage2_instance_partition_embed_weight"),
)


def worksheet(book: Workbook, title: str, headers: list[str], rows: list[dict]):
    sheet = book.create_sheet(title)
    sheet.append(headers)
    for item in rows:
        sheet.append([item.get(key) for key in headers])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        letter = column[0].column_letter
        width = min(58, max(13, max(len(str(cell.value or "")) for cell in column) + 2))
        sheet.column_dimensions[letter].width = width


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=ROOT / "review_methods")
    args = parser.parse_args()
    output = args.output.resolve()
    if output.exists():
        raise FileExistsError(f"Completed review protected: {output}")
    runs = []
    summary, per_plot, paired, config_rows = [], [], [], []
    for label, directory in RUNS:
        path = ROOT / directory / "result.json"
        result = json.loads(path.read_text())
        if len(result["per_plot"]) != 42:
            raise ValueError(f"Incomplete 14-plot comparison: {path}")
        methods = {item["method"]: item for item in result["comparison"]}
        fixed = methods["fixed"]
        for method in result["comparison"]:
            method = dict(run=label, **method)
            method["relative_compression_error"] = (
                method["mean_compression"] / fixed["mean_compression"] - 1)
            method["stage2_gate_pass"] = (method["method"] != "fixed"
                and abs(method["relative_compression_error"]) <= .05
                and method["mean_oracle_pq"] > fixed["mean_oracle_pq"]
                and method["mean_boundary_recall"] > fixed["mean_boundary_recall"])
            summary.append(method)
        per_plot.extend(dict(run=label, **item) for item in result["per_plot"])
        for method, metric in result["paired_vs_fixed"].items():
            for key, values in metric.items():
                paired.append(dict(run=label, method=method, metric=key,
                                   n=values["n"], mean_delta=values["mean_delta"],
                                   ci95_low=values["paired_bootstrap_ci95"][0],
                                   ci95_high=values["paired_bootstrap_ci95"][1]))
        config_rows.extend(dict(run=label, key=key, value=json.dumps(value))
                           for key, value in result["config"].items())
        runs.append(dict(run=label, path=str(path), val_plots=14,
                         full_model=False, test_used=False))
    training = json.loads((ROOT / "stage2_instance_embedding/train.json").read_text())
    output.mkdir(parents=True)
    review = dict(note="Partition kernels only; not full SPT/EZ-SP networks or end-to-end inference",
                  runs=runs, summary=summary, paired_vs_fixed=paired,
                  instance_embedding_checkpoint=training["checkpoint"],
                  instance_embedding_training=training["history"])
    (output / "review.json").write_text(json.dumps(review, indent=2) + "\n")
    workbook = Workbook()
    workbook.remove(workbook.active)
    worksheet(workbook, "runs", ["run", "path", "val_plots", "full_model", "test_used"], runs)
    worksheet(workbook, "summary", ["run", "method", "plots", "mean_compression",
        "relative_compression_error", "mean_oracle_pq", "mean_boundary_recall",
        "mean_mixed_group_fraction", "mean_known_point_purity",
        "mean_partition_seconds", "stage2_gate_pass"], summary)
    worksheet(workbook, "per_plot", ["run", "dataset_id", "source_dataset", "collection",
        "method", "points", "superpoints", "compression", "target_compression",
        "oracle_pq", "boundary_pairs", "boundary_recall", "mixed_group_fraction",
        "known_point_purity", "partition_seconds", "regularization",
        "peak_vram_mb"], per_plot)
    worksheet(workbook, "paired_bootstrap", ["run", "method", "metric", "n", "mean_delta",
        "ci95_low", "ci95_high"], paired)
    worksheet(workbook, "embedding_train", ["epoch", "updates", "mean_loss",
        "mean_positive_affinity", "mean_negative_affinity", "seconds"], training["history"])
    worksheet(workbook, "config", ["run", "key", "value"], config_rows)
    workbook.save(output / "review.xlsx")
    checked = load_workbook(output / "review.xlsx", read_only=True)
    assert len(checked.sheetnames) == 6
    checked.close()
    print(output / "review.xlsx")


if __name__ == "__main__":
    main()
