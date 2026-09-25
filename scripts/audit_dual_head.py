#!/usr/bin/env python3
"""Comparable validation point metrics for the legacy branch and selected masks."""
import argparse
import csv
import json
import sys
from pathlib import Path
from types import SimpleNamespace
import numpy as np
import torch
from openpyxl import load_workbook

PROJECT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT))
from pointcloud.model import LitePTTreeInstance
from pointcloud.data import load_npz, read_manifest
from pointcloud.instance_output import point_instance_metrics
from scripts.evaluate_pointcloud_litept import predict_plot, cluster_candidates, filter_candidates
from scripts.evaluate_combined_full_crowns import aggregate


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('--run', type=Path, default=PROJECT / 'outputs/dual_head_satv2_litept_v3')
    args = p.parse_args()
    torch.set_num_threads(4)
    torch.set_float32_matmul_precision('high')
    if not torch.cuda.is_available():
        raise RuntimeError('CUDA required')
    config = json.loads((args.run / 'configuration.json').read_text())
    selected = json.loads((args.run / 'selected.json').read_text())
    legacy = LitePTTreeInstance().to('cuda:0')
    old = torch.load(config['initial_weights'], map_location='cpu', weights_only=False)['model']
    legacy.load_state_dict(old)
    legacy.eval()
    legacy.backbone.shuffle_orders = False
    full = torch.load(args.run / 'weights/best.pt', map_location='cpu', weights_only=False)
    preserved = all(torch.equal(v, full['model']['legacy.'+k]) for k, v in old.items())
    if not preserved:
        raise AssertionError('The selected model changed legacy tensors')
    rows = []
    params = SimpleNamespace(tile_size=20., overlap=8., max_points=40000, preserve_height=True)
    for row in read_manifest(Path(config['manifest']), 'val'):
        arrays = load_npz(row['output'])
        raw = predict_plot(legacy, arrays, params, torch.device('cuda:0'), 20260925)
        candidates = filter_candidates(cluster_candidates(raw, config['legacy_cluster_config'], return_members=True), raw, config['legacy_cluster_config'])
        candidates = [c for c in candidates if c['height'] >= 2]
        labels = np.zeros(len(raw['coord']), np.uint32)
        for i, candidate in enumerate(candidates, 1):
            labels[candidate['point_indices']] = i
        rows.append({**{k: row[k] for k in ('dataset_id', 'source_dataset', 'collection', 'annotation_method')},
                     **point_instance_metrics(raw['gt_tree_id'], labels)})
        print(f'legacy point audit {len(rows)}: {row["dataset_id"]}', flush=True)
    new_rows = [{**{k: r[k] for k in ('dataset_id', 'source_dataset', 'collection', 'annotation_method')},
                 **r['point_metrics']} for r in selected['validation']['per_plot']]
    def summarize(values):
        return {**aggregate(values), 'mucov': float(np.mean([v['mucov'] for v in values])),
                'mwcov': float(np.mean([v['mwcov'] for v in values]))}
    report = dict(legacy_point_metrics=summarize(rows), new_point_metrics=summarize(new_rows),
                  legacy_native_point_metrics=summarize([r for r in rows if r['annotation_method']=='point_native']),
                  new_native_point_metrics=summarize([r for r in new_rows if r['annotation_method']=='point_native']),
                  legacy_tensors_bitwise_preserved=preserved, selected_epoch=selected['epoch'],
                  checkpoint_sha256=selected['checkpoint_sha256'], split='val', test_used=False,
                  caveat='Point metrics score all labelled-split voxels, including background; crown-ignore polygon protocol is separate.',
                  legacy_per_plot=rows)
    (args.run / 'point_comparison.json').write_text(json.dumps(report, indent=2)+'\n')
    wb = load_workbook(args.run / 'experiments.xlsx')
    if 'Porownanie walidacja' in wb:
        del wb['Porownanie walidacja']
    crown_sheet = wb.create_sheet('Porownanie walidacja')
    crown_sheet.append(['branch', 'source_balanced_crown_PQ', 'pooled_crown_F1', 'pooled_crown_PQ', 'checkpoint'])
    legacy_crowns = json.loads((args.run / 'legacy_validation/val_metrics.json').read_text())['metrics']
    for name, m, path in [('legacy', legacy_crowns, config['initial_weights']),
                           ('new_masks', selected['validation']['metrics'], str(args.run / 'weights/best.pt'))]:
        crown_sheet.append([name, m['source_balanced_pq'], m['f1'], m['pq'], path])
    if 'Metryki punktowe' in wb:
        del wb['Metryki punktowe']
    sheet = wb.create_sheet('Metryki punktowe')
    sheet.append(['branch_subset', 'source_balanced_PQ', 'pooled_PQ', 'F1', 'precision', 'recall', 'MUCov', 'MWCov'])
    for key in ('legacy_point_metrics', 'new_point_metrics', 'legacy_native_point_metrics', 'new_native_point_metrics'):
        m = report[key]
        sheet.append([key] + [m[k] for k in ('source_balanced_pq', 'pq', 'f1', 'precision', 'recall', 'mucov', 'mwcov')])
    sheet.freeze_panes='A2'
    if 'Wybrana konfiguracja' in wb:
        del wb['Wybrana konfiguracja']
    sheet = wb.create_sheet('Wybrana konfiguracja')
    sheet.append(['epoch', selected['epoch']])
    sheet.append(['checkpoint_sha256', selected['checkpoint_sha256']])
    for key, value in selected['config'].items():
        sheet.append([key, value])
    for sheet in wb:
        sheet.freeze_panes='A2'
    wb.save(args.run / 'experiments.xlsx')
    print(json.dumps({k: report[k] for k in ('legacy_point_metrics','new_point_metrics','selected_epoch')}, indent=2), flush=True)


if __name__ == '__main__':
    main()
