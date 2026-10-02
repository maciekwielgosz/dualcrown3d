#!/usr/bin/env python3
"""Summarize actual Stage-3 predictions and verify exported instance IDs."""
from __future__ import annotations
import argparse
from collections import defaultdict
import json
from pathlib import Path
import sys

import geopandas as gpd
import laspy
import numpy as np
import torch
from openpyxl import load_workbook

PROJECT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT))
from scripts.benchmark_superpoint_algorithms import save_json
from scripts.train_superpoint_decoder_pilot import DEFAULT, write_excel


def balanced(rows, metric):
    source = defaultdict(list)
    for row in rows:
        source[(row['source_dataset'], row['collection'])].append(row[metric])
    return np.mean([sum(r['iou_sum'] for r in values) /
        max(sum(r['tp'] + .5 * (r['fp'] + r['fn']) for r in values), 1.) for values in source.values()])


def paired_interval(first, second, metric, count=2000):
    before = {row['dataset_id']: row for row in first}
    after = {row['dataset_id']: row for row in second}
    if set(before) != set(after):
        raise ValueError('Unmatched validation plots')
    sources = defaultdict(list)
    for key, row in before.items():
        sources[(row['source_dataset'], row['collection'])].append(key)
    rng = np.random.default_rng(20261001)
    samples = []
    for _ in range(count):
        selected = [key for keys in sources.values() for key in rng.choice(keys, len(keys), replace=True)]
        samples.append(balanced([after[key] for key in selected], metric) -
                       balanced([before[key] for key in selected], metric))
    return dict(delta=float(balanced(second, metric) - balanced(first, metric)),
                lower=float(np.quantile(samples, .025)), upper=float(np.quantile(samples, .975)),
                procedure='paired within-source plot bootstrap; exploratory, one seed, selected validation checkpoints')


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('--output', type=Path, default=DEFAULT)
    args = p.parse_args()
    root = args.output.resolve()
    write_excel(root)
    selections = {}
    artifacts = []
    times = []
    for name in ('retained', 'control', 'fixed', 'ezsp'):
        run = root / 'runs' / name
        if not (run / 'DONE.json').exists():
            raise RuntimeError(f'Incomplete experiment: {name}')
        selection = json.loads((run / 'selected.json').read_text())
        selections[name] = selection
        config = json.loads((run / 'configuration.json').read_text())
        history = json.loads((run / 'history.json').read_text())['epochs']
        timing = dict(method=name, training_seconds=sum(row.get('seconds', 0) for row in history),
                      peak_training_vram_mb=max(row.get('peak_vram_mb', 0) for row in history),
                      mean_decoder_seconds=float(np.mean([row['decoder_seconds'] for row in selection['validation']['per_plot']])),
                      epoch=selection['epoch'], gpu=config['gpu'])
        if name in ('fixed', 'ezsp'):
            timing['groups_reassigned_within_superpoint'] = sum(row.get('groups_split_by_final_prediction', 0)
                                                               for row in selection['validation']['per_plot'])
            state = torch.load(selection['checkpoint'], map_location='cpu', weights_only=False)['model']
            timing['graph_output_weight_norm'] = float(state['decoder.graph.point_update.2.weight'].norm())
        times.append(timing)
        export = run / 'validation_exports'
        manifest = json.loads((export / 'export_manifest.json').read_text())
        for item in manifest['artifacts']:
            cloud = laspy.read(item['laz'])
            identifiers = set(np.unique(cloud.tree_id).tolist()) - {0}
            polygon = export / 'Segmentation3' / f"crowns_{item['dataset_id']}.gpkg"
            tops = export / 'Segmentation3' / f"ttops_{item['dataset_id']}.gpkg"
            if identifiers:
                frame = gpd.read_file(polygon)
                top_frame = gpd.read_file(tops)
                if set(frame.tree_id) != identifiers or set(top_frame.tree_id) != identifiers:
                    raise ValueError(f'LAZ/crown/treetop ID mismatch: {polygon}')
                if not frame.geometry.is_valid.all() or frame.geometry.is_empty.any():
                    raise ValueError(f'Invalid crown geometry: {polygon}')
                if any(not geom.covers(top) for geom, top in zip(frame.geometry, top_frame.geometry)):
                    raise ValueError(f'Treetop outside crown footprint: {polygon}')
            artifacts.append(dict(method=name, dataset_id=item['dataset_id'], instances=len(identifiers),
                                  points=len(cloud.points), matching_laz_gpkg_ids=True))
    intervals = {}
    for first, second in [('control', 'fixed'), ('control', 'ezsp'), ('fixed', 'ezsp')]:
        intervals[f'{second}_minus_{first}'] = {metric: paired_interval(
            selections[first]['validation']['per_plot'], selections[second]['validation']['per_plot'], metric)
            for metric in ('point', 'crown')}
    prepared = json.loads((root / 'prepared.json').read_text())
    val = [row for row in prepared['entries'] if row['split'] == 'val']
    preparation = {key: float(np.mean([row[key] for row in val])) for key in (
        'fixed_partition_seconds', 'fixed_graph_seconds', 'embedding_seconds',
        'ezsp_partition_seconds', 'ezsp_budget_search_seconds', 'ezsp_graph_seconds',
        'fixed_compression', 'ezsp_compression')}
    summary = json.loads((root / 'comparison.json').read_text())
    summary.update(paired_intervals=intervals, times=times, preparation=preparation,
                   artifacts_verified=artifacts, limitations=[
                       'One deterministic crop per plot, one seed, previously explored validation set.',
                       'Frozen encoder previously exposed to HELIOS; no backbone fine-tuning.',
                       'Retained reference is its mask decoder under shared pilot export, not deployed dual-consensus.',
                       'Crop-level point/crown metrics cannot be compared directly with full-scene production metrics.',
                       'Partition search and graph construction are preprocessing costs; decoder times are not full inference.',
                       'Independent validation and full-scene overlap fusion remain subsequent work.'])
    save_json(root / 'final_report.json', summary)
    book = load_workbook(root / 'experiments.xlsx')
    for name, records in [('timings', times), ('paired_intervals', [dict(comparison=key, metric=m, **v)
                          for key, metrics in intervals.items() for m, v in metrics.items()]),
                          ('export_checks', artifacts)]:
        if name in book:
            del book[name]
        sheet = book.create_sheet(name)
        fields = list(dict.fromkeys(key for row in records for key in row))
        sheet.append(fields)
        for row in records:
            sheet.append([row.get(key) for key in fields])
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
    book.save(root / 'experiments.xlsx')
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    names = [row['method'] for row in summary['comparison']]
    x = np.arange(len(names))
    figure, axes = plt.subplots(1, 2, figsize=(10, 4), constrained_layout=True)
    for axis, metric, title in zip(axes, ['point_pq', 'crown_pq'], ['Point instance PQ', 'Crown footprint PQ']):
        values = [row[metric] for row in summary['comparison']]
        axis.bar(x, values, color=['#527994', '#3c9d92', '#bd7737', '#7862a2'])
        axis.set_xticks(x, names)
        axis.set_ylim(0, min(1., max(values) * 1.3))
        axis.set_title(title)
        axis.bar_label(axis.containers[0], fmt='%.3f', padding=3)
        axis.set_ylabel('Source-balanced PQ @ IoU 0.5')
    figure.suptitle('Stage-3 pilot: 14 validation crops, selected checkpoints; one seed')
    figure.savefig(root / 'comparison.png', dpi=160)
    plt.close(figure)
    print(json.dumps({key: summary[key] for key in ('comparison', 'paired_intervals', 'times', 'preparation')}, indent=2))


if __name__ == '__main__':
    main()
