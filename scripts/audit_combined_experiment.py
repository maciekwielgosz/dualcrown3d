#!/usr/bin/env python3
"""Check delivered polygons and preserve reproducibility metadata after final test."""
import csv
import hashlib
import importlib.metadata
import json
import os
import shutil
import sys
from datetime import datetime, timezone
from pathlib import Path

import geopandas as gpd
import numpy as np
import torch
from openpyxl import load_workbook

PROJECT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT))
from pointcloud.experiment_log import ROOT, record
from scripts.evaluate_combined_full_crowns import PROTOCOL_VERSION

DATA_ROOT = Path(os.environ.get("SEGMENTATION_DATA_ROOT", PROJECT.parent / "combined_als_crowns_v1")).resolve()


def digest(path):
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(8*1024**2), b""):
            h.update(block)
    return h.hexdigest()


def main():
    selected = ROOT / "selected"
    frozen = json.loads((selected / "frozen_selection.json").read_text())
    with (DATA_ROOT / "manifest_test.csv").open() as stream:
        rows = list(csv.DictReader(stream))
    with (DATA_ROOT / "manifest.csv").open() as stream:
        all_rows = list(csv.DictReader(stream))
    excluded_idtrees = all(r["collection"] != "IDTREES" for r in all_rows)
    if (DATA_ROOT / "exclusion_report.json").exists():
        assert excluded_idtrees, "Filtered dataset still contains IDTREES"
    reports = []
    for method in ("model_test", "classical_test"):
        count, holes, multipart = 0, 0, 0
        for row in rows:
            folder = selected / method / "Segmentation3"
            crowns = gpd.read_file(folder / f"crowns_{row['dataset_id']}.gpkg")
            tops = gpd.read_file(folder / f"ttops_{row['dataset_id']}.gpkg")
            gt = gpd.read_file(row["gt_vector"])
            assert crowns.crs == gt.crs == tops.crs, row["dataset_id"]
            assert crowns.treeID.is_unique and tops.treeID.is_unique
            assert set(crowns.treeID) == set(tops.treeID)
            assert crowns.geometry.is_valid.all() and (~crowns.geometry.is_empty).all()
            assert tops.geometry.is_valid.all() and tops.geometry.has_z.all()
            assert np.isfinite(tops.Z).all()
            assert np.allclose(crowns.area_m2, crowns.geometry.area)
            for geometry in crowns.geometry:
                assert geometry.geom_type in ("Polygon", "MultiPolygon")
                parts = list(geometry.geoms) if geometry.geom_type == "MultiPolygon" else [geometry]
                holes += sum(len(p.interiors) for p in parts)
                multipart += len(parts) > 1
                if method == "model_test":
                    assert len(parts) == 1 and not parts[0].interiors, row["dataset_id"]
            count += len(crowns)
        assert len(list(folder.glob("crowns_*.gpkg"))) == len(rows)
        assert len(list(folder.glob("ttops_*.gpkg"))) == len(rows)
        reports.append(dict(method=method, plot_pairs=len(rows), crowns=count, holes=holes, multipart=multipart))
    # Preserve earlier snapshots; source_final is the code state used for the
    # delivered report and post-test artifact audit.
    archive = ROOT / "reproducibility/source_final_v2"
    sources = list((PROJECT / "pointcloud").glob("*.py"))
    sources += [p for p in (PROJECT / "scripts").glob("*.py") if any(word in p.stem for word in ("combined", "pointcloud", "offset_checkpoints"))]
    sources += list((PROJECT / "configs").glob("combined*_campaign.json"))
    hashes = {}
    for path in sources:
        destination = archive / path.relative_to(PROJECT)
        destination.parent.mkdir(parents=True, exist_ok=True)
        checksum = digest(path)
        if destination.exists() and digest(destination) != checksum:
            raise RuntimeError(f"Refusing to replace an existing source snapshot: {destination}")
        if not destination.exists():
            shutil.copy2(path, destination)
        hashes[str(path.relative_to(PROJECT))] = checksum
    # Include the actual imported backbone and classical baseline code, not
    # just wrappers. No datasets, weights or compiled extension binaries here.
    dependencies = [(p, Path("vendor/LitePT") / p.relative_to(PROJECT / "vendor/LitePT"))
                    for p in (PROJECT / "vendor/LitePT").rglob("*.py")]
    dependencies += [(PROJECT.parent / "run_r/code" / name, Path("run_r_reference") / name)
                     for name in ("pcopw_chunks_500m_Segmentacja.py", "for_instance_to_chm_gt.py", "evaluate_structural_ensemble.py")]
    for path, relative in dependencies:
        destination = archive / relative
        destination.parent.mkdir(parents=True, exist_ok=True)
        checksum = digest(path)
        if destination.exists() and digest(destination) != checksum:
            raise RuntimeError(f"Dependency source snapshot changed: {destination}")
        if not destination.exists():
            shutil.copy2(path, destination)
        hashes[str(relative)] = checksum
    packages = {d.metadata["Name"]: d.version for d in importlib.metadata.distributions()}
    (ROOT / "reproducibility/environment.json").write_text(json.dumps(dict(python=sys.version, executable=sys.executable, packages=packages, torch_cuda=torch.version.cuda), indent=2)+"\n")
    (ROOT / "reproducibility/source_sha256.json").write_text(json.dumps(hashes, indent=2)+"\n")
    # Rebuild workbook from current protocol, marking provisional historical rows.
    for path in (ROOT / "records").glob("*.json"):
        data = json.loads(path.read_text())
        name = data["experiment_id"]
        if name.startswith("classical"):
            result_path = ROOT / name / f"val_metrics_{PROTOCOL_VERSION}.json"
            if result_path.exists():
                result = json.loads(result_path.read_text())
                record(name, final_val_metrics=result["metrics"], protocol_version=PROTOCOL_VERSION)
        elif "training_dir" in data and not name.startswith("smoke"):
            result_path = ROOT / name / "final_validation/selection.json"
            if result_path.exists():
                result = json.loads(result_path.read_text())
                assert result["protocol_version"] == PROTOCOL_VERSION
                record(name, protocol_version=PROTOCOL_VERSION, final_val_metrics=result["metrics"],
                       early_validation_caveat="Early epoch/selected_val scores preceded grid and annotation-ignore fixes; compare final_val and test only.")
    workbook = load_workbook(ROOT / "experiments.xlsx", read_only=True)
    assert {"Eksperymenty", "Parametry", "Metryki_zrodla", "Epoki", "Podzial", "Historyczne_run_r", "Protokol"}.issubset(workbook.sheetnames)
    report = dict(created_utc=datetime.now(timezone.utc).isoformat(), protocol_version=PROTOCOL_VERSION,
                  dataset_manifest_sha256=frozen["dataset_manifest_sha256"], outputs=reports,
                  dataset_plots=len(all_rows), idtrees_absent=excluded_idtrees,
                  model_checkpoint_sha256=digest(selected / "best.pt"), workbook_sheets=workbook.sheetnames,
                  workbook_experiment_rows=workbook["Eksperymenty"].max_row-1, source_files=len(hashes))
    workbook.close()
    (selected / "artifact_validation.json").write_text(json.dumps(report, indent=2)+"\n")
    print(json.dumps(report, indent=2))


if __name__ == "__main__":
    main()
