#!/usr/bin/env python3
"""Train the second-pass refiner on frozen Model20 (stage-1) predictions.

Selection uses all native validation plots and the gate preregistered in
protocol.json before training. The held-out test split is never read here.
"""
import argparse
import csv
import itertools
import json
import math
import os
from collections import Counter
from pathlib import Path
import sys
import time

import geopandas as gpd
import numpy as np
import shapely
import torch
from openpyxl import Workbook
from torch.utils.data import DataLoader, Dataset

PROJECT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT))
from pointcloud.data import load_npz, move_to_device
from pointcloud.dual_fusion import instance_records
from pointcloud.instance_output import point_instance_metrics
from pointcloud.joint_training import fixed_serialization
from pointcloud.two_pass import (CONDITION_DIM, RefinementDecoder, corrupt_stage1, crop_with_extras, need_target,
                                 per_point_matched_label, refinement_losses, seed_refinement_losses,
                                 stage1_conditioning)
from pointcloud.two_pass_inference import STAGE1_KEYS, frozen_window, stage2_sweep
from pointcloud.two_pass_reconcile import DEFAULT_RECONCILE, reconcile, segmentation_diagnostics
from scripts.cache_stage1_predictions import DEFAULT_OUTPUT, eligible_rows, sha256, write_json
from scripts.calibrate_legacy_small_trees import small_hits
from scripts.evaluate_combined_full_crowns import aggregate, metrics
from scripts.train_supervision_v4 import build

NAMES = ('dataset_id', 'source_dataset', 'collection')


def load_stage1(path):
    with np.load(path) as archive:
        stage1 = {k: archive[k] for k in STAGE1_KEYS}
        return stage1, archive['missed_gt_ids'], archive['missed_gt_counts']


class RefinementCrops(Dataset):
    """Crops of real train plots with their (optionally corrupted) stage-1 result."""
    def __init__(self, rows, cache, seed, max_points, absorb, drop):
        if any(r['model_split'] != 'train' for r in rows):
            raise ValueError('RefinementCrops accepts train rows only')
        self.rows, self.cache, self.seed, self.max_points = rows, Path(cache), seed, max_points
        self.absorb, self.drop, self.epoch = absorb, drop, 0

    def __len__(self):
        return len(self.rows)

    def __getitem__(self, index):
        row_index, draw = index
        row = self.rows[row_index]
        rng = np.random.default_rng(np.random.SeedSequence([self.seed, self.epoch, draw, row_index]))
        arrays = load_npz(row['output'])
        stage1, missed, missed_counts = load_stage1(self.cache / f"{row['dataset_id']}.npz")
        ids = arrays['tree_id']
        anchor, choice = None, rng.random()
        if choice < .5 and len(missed):
            weight = 1. / np.sqrt(np.maximum(missed_counts, 1))
            target = rng.choice(missed, p=weight / weight.sum())
            anchor = int(rng.choice(np.flatnonzero(ids == target)))
        elif choice >= .8:
            anchor = int(rng.integers(len(ids)))
        extras = dict(stage1_labels=stage1['labels'], stage1_confidence=stage1['confidence'],
                      stage1_source=stage1['source'], tree_probability=stage1['tree_probability'],
                      point_probability=stage1['point_probability'],
                      candidate_count=stage1['candidate_count'], foreign_support=stage1['foreign_support'])
        crop = crop_with_extras(arrays, extras, dict(vote_xy=stage1['vote_xy']), rng, 20., self.max_points,
                                True, True, anchor)
        xyz = crop['coord'].numpy()
        labels, confidence, source = corrupt_stage1(xyz, crop['stage1_labels'].numpy(),
                                                    crop['stage1_confidence'].numpy(),
                                                    crop['stage1_source'].numpy(), rng, self.absorb, self.drop)
        condition = stage1_conditioning(xyz, crop['vote_xy'].numpy(), labels, confidence, source,
                                        crop['tree_probability'].numpy(), crop['point_probability'].numpy(),
                                        crop['candidate_count'].numpy(), crop['foreign_support'].numpy())
        tree_id = crop['tree_id'].numpy()
        matched = per_point_matched_label(tree_id, labels)
        return dict(coord=crop['coord'], grid_coord=crop['grid_coord'], feat=crop['feat'], offset=crop['offset'],
                    tree_id=crop['tree_id'], condition=torch.from_numpy(condition),
                    gt_matched_label=torch.from_numpy(matched), stage1_labels=torch.from_numpy(labels),
                    need_target=torch.from_numpy(need_target(tree_id, labels, matched)))


def balanced_weights(rows):
    groups = [(r['source_dataset'], r['collection']) for r in rows]
    counts = Counter(groups)
    return np.asarray([1. / len(counts) / counts[g] for g in groups], np.float64)


def reconcile_sweep():
    sizes = {'m12': dict(minimum_voxels=12, minimum_height_m=2., minimum_area_m2=.75),
             'm8': dict(minimum_voxels=8, minimum_height_m=1.5, minimum_area_m2=.4)}
    configs = {}
    for p, q, (name, size) in itertools.product((.3, .5, .7), (.2, .4), sizes.items()):
        configs[f'new{p:g}_q{q:g}_{name}'] = {**DEFAULT_RECONCILE, 'new_probability': p, 'min_quality': q, **size}
    base = configs['new0.5_q0.4_m12']
    configs['new0.5_q0.4_m12_merge'] = {**base, 'allow_merge': True}
    configs['new0.5_q0.4_m12_completion'] = {**base, 'apply_completion': True}
    configs['new0.5_q0.4_m12_take0.7'] = {**base, 'max_take_fraction': .7}
    configs['new0.5_q0.2_m8_take0.7'] = {**configs['new0.5_q0.2_m8'], 'max_take_fraction': .7}
    return configs


def evaluate_plot(row, arrays, gt, ignore, labels, confidence):
    geometries = [r['geometry'] for r in instance_records(arrays, labels, confidence)]
    return dict(**{k: row[k] for k in NAMES}, **metrics(gt.geometry, geometries, ignore),
                point_metrics=point_instance_metrics(arrays['tree_id'], labels),
                **small_hits(gt.geometry, geometries),
                diagnostics=segmentation_diagnostics(arrays['tree_id'], labels),
                predicted_instances=len(geometries))


def summarize(records):
    crown = aggregate(records)
    point = aggregate([{**{k: r[k] for k in NAMES}, **r['point_metrics']} for r in records])
    total = lambda key: int(sum(r[key] for r in records))
    diag = lambda key: int(sum(r['diagnostics'][key] for r in records))
    by_source = {}
    for r in records:
        entry = by_source.setdefault(r['collection'], dict(small_10_tp=0, small_10_gt=0))
        entry['small_10_tp'] += r['small_10_tp']
        entry['small_10_gt'] += r['small_10_gt']
    return dict(point_sb_pq=point['source_balanced_pq'], crown_sb_pq=crown['source_balanced_pq'],
                point_sb_f1=point['source_balanced_f1'], crown_sb_f1=crown['source_balanced_f1'],
                point_precision=point['precision'], crown_precision=crown['precision'],
                point_recall=point['recall'], crown_recall=crown['recall'],
                small_10_tp=total('small_10_tp'), small_10_gt=total('small_10_gt'),
                small_4_tp=total('small_4_tp'), small_4_gt=total('small_4_gt'),
                large_recall=diag('large_tp') / max(diag('large_gt'), 1),
                sparse_point_tp=diag('small_tp'), sparse_point_gt=diag('small_gt'),
                oversplit_gt=diag('oversplit_gt'), undersegmented_pred=diag('undersegmented_pred'),
                predicted_instances=total('predicted_instances'), small_by_collection=by_source)


def gate_checks(summary, base, gate):
    return dict(
        point_sb_pq=summary['point_sb_pq'] >= base['point_sb_pq'] - gate['point_sb_pq_tolerance'],
        crown_sb_pq=summary['crown_sb_pq'] >= base['crown_sb_pq'] - gate['crown_sb_pq_tolerance'],
        point_precision=summary['point_precision'] >= base['point_precision'] - gate['precision_tolerance'],
        crown_precision=summary['crown_precision'] >= base['crown_precision'] - gate['precision_tolerance'],
        large_recall=summary['large_recall'] >= base['large_recall'] - gate['large_recall_tolerance'],
        small_crowns=summary['small_10_tp'] >= gate['small_10_minimum_tp'])


def rank(summary):
    return (summary['small_10_tp'], math.sqrt(max(summary['point_sb_pq'] * summary['crown_sb_pq'], 0.)))


def validate(model, refiner, rows, cache, configs, context):
    records = {name: [] for name in configs}
    seconds = dict(stage2=0., reconcile=0.)
    proposals_total = 0
    for row in rows:
        if row['dataset_id'] not in context:
            arrays = load_npz(row['output'])
            stage1, _, _ = load_stage1(cache / f"{row['dataset_id']}.npz")
            gt = gpd.read_file(row['gt_vector'])
            ignore = shapely.from_wkb(Path(row['ignore_vector']).read_bytes()) if row.get('ignore_vector') else None
            context[row['dataset_id']] = (arrays, stage1, gt, ignore)
        arrays, stage1, gt, ignore = context[row['dataset_id']]
        proposals = stage2_sweep(model, refiner, arrays, stage1, max_points=40000)
        seconds['stage2'] += float(proposals['seconds'])
        proposals_total += len(proposals['quality'])
        for name, config in configs.items():
            start = time.monotonic()
            labels, confidence, _, operations = reconcile(arrays, stage1['labels'], stage1['confidence'],
                                                          stage1['source'], proposals, config)
            seconds['reconcile'] += time.monotonic() - start
            record = evaluate_plot(row, arrays, gt, ignore, labels, confidence)
            record['operations'] = dict(Counter(op['op'] for op in operations))
            record['taken_voxels'] = int(sum(op.get('taken', 0) for op in operations))
            records[name].append(record)
    return ({name: summarize(value) for name, value in records.items()}, records, seconds, proposals_total)


def stage1_baseline(rows, cache, context, protocol):
    records = []
    for row in rows:
        arrays = load_npz(row['output'])
        stage1, _, _ = load_stage1(cache / f"{row['dataset_id']}.npz")
        gt = gpd.read_file(row['gt_vector'])
        ignore = shapely.from_wkb(Path(row['ignore_vector']).read_bytes()) if row.get('ignore_vector') else None
        context[row['dataset_id']] = (arrays, stage1, gt, ignore)
        records.append(evaluate_plot(row, arrays, gt, ignore, stage1['labels'].astype(np.uint32), stage1['confidence']))
    summary = summarize(records)
    if len(rows) == protocol['validation_plots']:
        for key in ('point_sb_pq', 'crown_sb_pq', 'point_sb_f1', 'crown_sb_f1'):
            if abs(summary[key] - protocol['baseline'][key]) > 1e-6:
                raise AssertionError(f'Frozen Model20 baseline not reproduced: {key} '
                                     f'{summary[key]} vs {protocol["baseline"][key]}')
    return summary, records


def write_excel(run, history, validations, selection):
    book = Workbook()
    sheet = book.active
    sheet.title = 'epochs'
    fields = list(dict.fromkeys(k for r in history for k in r))
    if fields:
        sheet.append(fields)
        for r in history:
            sheet.append([r.get(k) for k in fields])
    tab = book.create_sheet('validation')
    keys = ('point_sb_pq', 'crown_sb_pq', 'point_sb_f1', 'crown_sb_f1', 'point_precision', 'crown_precision',
            'small_10_tp', 'small_10_gt', 'small_4_tp', 'small_4_gt', 'large_recall', 'oversplit_gt',
            'undersegmented_pred', 'predicted_instances')
    tab.append(('epoch', 'config', 'gate_passed', *keys))
    for epoch, name, passed, summary in validations:
        tab.append((epoch, name, passed, *[summary[k] for k in keys]))
    tab.freeze_panes = 'A2'
    final = book.create_sheet('selection')
    for key, value in selection.items():
        final.append((key, json.dumps(value, default=str) if isinstance(value, (dict, list)) else value))
    temporary = run / 'experiments.tmp.xlsx'
    book.save(temporary)
    os.replace(temporary, run / 'experiments.xlsx')


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument('--run', default='residual_v1')
    parser.add_argument('--epochs', type=int, default=30)
    parser.add_argument('--samples', type=int, default=240)
    parser.add_argument('--val-every', type=int, default=3)
    parser.add_argument('--seed', type=int, default=20261002)
    parser.add_argument('--queries', type=int, default=64)
    parser.add_argument('--layers', type=int, default=3)
    parser.add_argument('--memory-tokens', type=int, default=1024)
    parser.add_argument('--lr', type=float, default=3e-4)
    parser.add_argument('--max-points', type=int, default=16000)
    parser.add_argument('--absorb', type=float, default=.3)
    parser.add_argument('--drop', type=float, default=.15)
    parser.add_argument('--missed-weight', type=float, default=2.)
    parser.add_argument('--teacher-epochs', type=int, default=8)
    parser.add_argument('--targets', choices=('seed', 'hungarian'), default='seed',
                        help='seed: v2 containment targets; hungarian: v1 set matching')
    parser.add_argument('--no-context-priors', action='store_true', help='v1 decoder without instance/lid priors')
    parser.add_argument('--center-shift', action='store_true', help='v3: learned per-query centre shift')
    parser.add_argument('--accumulation', type=int, default=2)
    parser.add_argument('--warm-start', action='store_true',
                        help="v3: start from Model20's mask decoder with a zero-initialised conditioned adapter")
    parser.add_argument('--copied-lr', type=float, default=1e-4, help='learning rate of copied Model20 modules')
    parser.add_argument('--smoke', action='store_true', help='Few steps and one validation plot; nothing is selected')
    args = parser.parse_args()
    torch.set_num_threads(4)
    if not torch.cuda.is_available():
        raise RuntimeError('GPU is required')
    device = torch.device('cuda:0')
    root = args.output.resolve()
    protocol = json.loads((root / 'protocol.json').read_text())
    if sha256(protocol['stage1_checkpoint']) != protocol['stage1_sha256'] or sha256(protocol['manifest']) != protocol['manifest_sha256']:
        raise ValueError('Stage-1 checkpoint or manifest changed since the protocol was frozen')
    train_rows, val_rows = eligible_rows('train'), eligible_rows('val')
    for split, rows in (('train', train_rows), ('val', val_rows)):
        missing = [r['dataset_id'] for r in rows if not (root / 'stage1' / split / f"{r['dataset_id']}.npz").exists()]
        if missing:
            raise FileNotFoundError(f'Run cache_stage1_predictions.py first; missing {split}: {missing[:3]}')
    run = root / 'runs' / (args.run + ('_smoke' if args.smoke else ''))
    if (run / 'DONE.json').exists():
        print((run / 'selected.json').read_text())
        return
    if run.exists() and not args.smoke:
        raise FileExistsError(f'Incomplete run protected: {run}')
    (run / 'weights').mkdir(parents=True, exist_ok=True)
    torch.manual_seed(args.seed)
    np.random.seed(args.seed)
    model, _ = build(protocol['stage1_checkpoint'])
    model.requires_grad_(False)
    model.eval()
    fixed_serialization(model)
    refiner_args = dict(queries=args.queries, layers=args.layers, memory_tokens=args.memory_tokens,
                        context_priors=not args.no_context_priors, condition_dim=CONDITION_DIM,
                        center_shift=args.center_shift, warm_start=args.warm_start)
    loss_function = seed_refinement_losses if args.targets == 'seed' else refinement_losses
    refiner = RefinementDecoder(**refiner_args).to(device)
    parameters = list(refiner.parameters())
    if args.warm_start:
        refiner.load_model20_decoder(model.point_decoder)
        copied = [p for name in refiner.COPIED for p in getattr(refiner, name).parameters()]
        identifiers = {id(p) for p in copied}
        fresh = [p for p in parameters if id(p) not in identifiers]
        optimizer = torch.optim.AdamW([dict(params=fresh, lr=args.lr), dict(params=copied, lr=args.copied_lr)],
                                      weight_decay=.01)
    else:
        optimizer = torch.optim.AdamW(parameters, lr=args.lr, weight_decay=.01)
    scheduler = torch.optim.lr_scheduler.CosineAnnealingLR(optimizer, T_max=args.epochs, eta_min=1e-6)
    configs = reconcile_sweep()
    if args.smoke:
        val_rows, configs = val_rows[:1], dict(list(configs.items())[:2])
        args.epochs, args.samples, args.val_every = 1, 6, 1
    configuration = dict(run=run.name, seed=args.seed, refiner_args=refiner_args,
                         trainable_parameters=sum(p.numel() for p in parameters),
                         architecture='frozen Model20 (LitePT-S + dense decoder) -> conditioned need head + '
                                      f'{args.queries} need/coverage-seeded queries, {args.layers} mask-attention '
                                      'layers, heads: mask / decision(background,new,existing) / quality',
                         loss='need BCE + mask(dice+focal, weighted missed trees) + decision CE + 0.5 quality MSE',
                         targets=args.targets,
                         sampling='source-balanced real train plots; anchors 50% stage-1-missed tree, 30% any tree, 20% any point',
                         augmentation=dict(rotation=True, scale='0.9-1.1', density=(.5, .75, 1.),
                                           stage1_absorb=args.absorb, stage1_drop=args.drop),
                         epochs=args.epochs, samples_per_epoch=args.samples, max_points_train=args.max_points,
                         lr=args.lr, copied_lr=args.copied_lr if args.warm_start else None, weight_decay=.01, accumulation=args.accumulation, teacher_epochs=args.teacher_epochs,
                         missed_weight=args.missed_weight, train_plots=len(train_rows), val_plots=len(val_rows),
                         protocol_sha256=sha256(root / 'protocol.json'), reconcile_configs=configs,
                         stage1_training_predictions='in-sample Model20 predictions (not out-of-fold); '
                                                     'mitigated by stage-1 corruption augmentation',
                         test_used_for_selection=False)
    write_json(run / 'configuration.json', configuration)
    context = {}
    base, base_records = stage1_baseline(val_rows, root / 'stage1/val', context, protocol)
    write_json(run / 'validation/stage1_only.json', dict(summary=base, per_plot=base_records))
    print(json.dumps(dict(stage1_only={k: base[k] for k in ('point_sb_pq', 'crown_sb_pq', 'small_10_tp', 'large_recall')})), flush=True)
    dataset = RefinementCrops(train_rows, root / 'stage1/train', args.seed, args.max_points, args.absorb, args.drop)
    weights = balanced_weights(train_rows)
    history, validations = [], []
    selected = None          # full gate
    preserving = None        # all non-small gates, best small-crown gain
    started = time.monotonic()

    def save(path, epoch, extra=None):
        temporary = path.with_suffix('.tmp.pt')
        torch.save(dict(refiner=refiner.state_dict(), refiner_args=refiner_args, epoch=epoch,
                        stage1_checkpoint=protocol['stage1_checkpoint'], stage1_sha256=protocol['stage1_sha256'],
                        stage1_config=protocol['stage1_config'], configuration=configuration, **(extra or {})), temporary)
        os.replace(temporary, path)

    for epoch in range(1, args.epochs + 1):
        refiner.train()
        model.eval()
        refiner.teacher_probability = max(0., 1. - (epoch - 1) / max(args.teacher_epochs, 1))
        dataset.epoch = epoch
        rng = np.random.default_rng(args.seed + epoch * 100003)
        draws = [(int(i), step) for step, i in enumerate(rng.choice(len(train_rows), args.samples, p=weights / weights.sum()))]
        loader = DataLoader(dataset, batch_size=None, sampler=draws, num_workers=2, pin_memory=True)
        totals = Counter()
        optimizer.zero_grad(set_to_none=True)
        start = time.monotonic()
        torch.cuda.reset_peak_memory_stats()
        for step, batch in enumerate(loader, 1):
            batch = move_to_device(batch, device)
            inputs = {k: batch[k] for k in ('coord', 'grid_coord', 'feat', 'offset')}
            features, dense = frozen_window(model, inputs)
            output = refiner(features, dense['features'], inputs, batch['condition'],
                             dense['semantic_logits'], dense['offset_m'], teacher_need=batch['need_target'],
                             stage1_labels=batch['stage1_labels'])
            losses = loss_function(output, batch, missed_weight=args.missed_weight)
            if not torch.isfinite(losses['loss']):
                raise FloatingPointError('Non-finite loss')
            (losses['loss'] / args.accumulation).backward()
            if step % args.accumulation == 0 or step == len(draws):
                torch.nn.utils.clip_grad_norm_(parameters, 2., error_if_nonfinite=True)
                optimizer.step()
                optimizer.zero_grad(set_to_none=True)
            totals.update({k: float(v.detach()) for k, v in losses.items()})
        scheduler.step()
        record = dict(epoch=epoch, seconds=time.monotonic() - start,
                      peak_vram_mb=torch.cuda.max_memory_allocated() / 1024 ** 2,
                      teacher_probability=refiner.teacher_probability,
                      **{f'train_{k}': v / len(draws) for k, v in totals.items()})
        if epoch % args.val_every == 0 or epoch == args.epochs:
            summaries, records, seconds, proposals = validate(model, refiner, val_rows, root / 'stage1/val', configs, context)
            write_json(run / f'validation/epoch_{epoch:03d}/metrics.json',
                       dict(summaries=summaries, per_plot=records, seconds=seconds, proposals=proposals))
            checks = {name: gate_checks(s, base, protocol['gate']) for name, s in summaries.items()}
            for name, summary in summaries.items():
                validations.append((epoch, name, all(checks[name].values()), summary))
            passing = [n for n in summaries if all(checks[n].values())]
            keeping = [n for n in summaries if all(v for k, v in checks[n].items() if k != 'small_crowns')]
            best_name = max(passing or keeping or summaries, key=lambda n: rank(summaries[n]))
            best = summaries[best_name]
            record.update(val_config=best_name, val_gate_passed=bool(passing), val_point_sb_pq=best['point_sb_pq'],
                          val_crown_sb_pq=best['crown_sb_pq'], val_small_10_tp=best['small_10_tp'],
                          val_large_recall=best['large_recall'], val_predicted=best['predicted_instances'],
                          val_stage2_seconds=seconds['stage2'], val_proposals=proposals)
            if passing and not args.smoke:
                name = max(passing, key=lambda n: rank(summaries[n]))
                if selected is None or rank(summaries[name]) > rank(selected['validation']):
                    selected = dict(epoch=epoch, config_name=name, config=configs[name], validation=summaries[name],
                                    checks=checks[name])
                    save(run / 'weights/best.pt', epoch, dict(reconcile_config=configs[name]))
            if keeping and not args.smoke:
                name = max(keeping, key=lambda n: rank(summaries[n]))
                if (summaries[name]['small_10_tp'] > base['small_10_tp']
                        and (preserving is None or rank(summaries[name]) > rank(preserving['validation']))):
                    preserving = dict(epoch=epoch, config_name=name, config=configs[name],
                                      validation=summaries[name], checks=checks[name])
                    save(run / 'weights/best_quality_preserving.pt', epoch, dict(reconcile_config=configs[name]))
        save(run / 'weights/last.pt', epoch)
        history.append(record)
        with (run / 'training_log.csv').open('w', newline='') as stream:
            fields = list(dict.fromkeys(k for r in history for k in r))
            writer = csv.DictWriter(stream, fieldnames=fields)
            writer.writeheader()
            writer.writerows(history)
        selection = dict(selected=selected, best_quality_preserving=preserving, stage1_only=base,
                         gate=protocol['gate'], test_used_for_selection=False)
        write_json(run / 'selected.json', selection)
        write_excel(run, history, validations, selection)
        print(json.dumps(record), flush=True)
    if not args.smoke:
        write_json(run / 'DONE.json', dict(epochs=args.epochs, seconds=time.monotonic() - started,
                                           selected_epoch=selected and selected['epoch'],
                                           quality_preserving_epoch=preserving and preserving['epoch']))
    print(json.dumps(dict(selected=selected and {k: selected[k] for k in ('epoch', 'config_name')},
                          best_quality_preserving=preserving and {k: preserving[k] for k in ('epoch', 'config_name')}),
                     indent=2), flush=True)


if __name__ == '__main__':
    main()
