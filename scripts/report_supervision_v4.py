#!/usr/bin/env python3
"""Reviewable v4 reports, weak-2D audit, paired timing and bootstrap intervals."""
import argparse
from collections import defaultdict
import csv
import json
from pathlib import Path
import sys
import time

import geopandas as gpd
import numpy as np
import torch
from openpyxl import load_workbook

PROJECT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(PROJECT))
from scripts.train_supervision_v4 import DEFAULT_OUTPUT,REAL,INITIAL,build,dump,configurations,excel,evaluate,digest
from pointcloud.data import read_manifest,load_npz
from pointcloud.instance_output import merge_masks_with_sources
from scripts.predict_dual_head import predict
from scripts.evaluate_combined_full_crowns import metrics,aggregate
from scripts.prepare_supervision_v4 import write_csv


def bootstrap(before,after,key,repeats=2000):
    left={r['dataset_id']:r for r in before['per_plot']};right={r['dataset_id']:r for r in after['per_plot']}
    if set(left)!=set(right):raise ValueError('Unpaired plot sets')
    grouped=defaultdict(list)
    for name,r in left.items():grouped[(r['source_dataset'],r['collection'])].append(name)
    def score(rows):
        total=sum(r['tp']+.5*(r['fp']+r['fn']) for r in rows)
        return sum(r['iou_sum'] for r in rows)/max(total,1.)
    def items(index,names):return [index[n][key] if key=='point_metrics' else index[n] for n in names]
    rng=np.random.default_rng(20261003);values=[]
    for _ in range(repeats):
        differences=[]
        for names in grouped.values():
            draw=rng.choice(names,len(names)).tolist()
            differences.append(score(items(right,draw))-score(items(left,draw)))
        values.append(np.mean(differences))
    delta=np.mean([score(items(right,n))-score(items(left,n)) for n in grouped.values()])
    return dict(delta=float(delta),lower_95=float(np.percentile(values,2.5)),upper_95=float(np.percentile(values,97.5)),
                repeats=repeats,caveat='Paired plots stratified by collection; conditional on one selected training seed. Small strata limit uncertainty.')


def architecture_ablation(root,cache_only=False):
    """Report the validation-frozen hybrid even if the baseline wins selection.

    These test results are descriptive ablations, never input to selection.
    """
    run=root/'runs/hybrid_queries'
    frozen=json.loads((run/'selected.json').read_text())
    path=run/'weights/best.pt'
    target=root/'final_test/hybrid_queries/real/metrics.json'
    if target.exists():result=json.loads(target.read_text())['fixed']
    else:
        if cache_only:raise FileNotFoundError(f'Run GPU evaluation before --skip-gpu: {target}')
        model,_=build(path)
        result=evaluate(model,REAL,'test',target.parent,{'fixed':frozen['config']})['fixed']
        del model;torch.cuda.empty_cache()
    summary=dict(model='hybrid_queries',domain='real',epoch=frozen['epoch'],fusion=frozen['fusion'],
        checkpoint=str(path),checkpoint_sha256=digest(path),test_used_for_selection=False,
        point_SB_PQ=result['point_metrics']['source_balanced_pq'],crown_SB_PQ=result['metrics']['source_balanced_pq'],
        point_F1=result['point_metrics']['f1'],crown_F1=result['metrics']['f1'])
    dump(root/'architecture_ablation.json',summary)
    return summary,result


def checkpoint_audit(root):
    original=torch.load(INITIAL,map_location='cpu',weights_only=False)['model']
    rows=[]
    for run in ('data_only','hybrid_queries'):
        for kind in ('best','last'):
            path=root/f'runs/{run}/weights/{kind}.pt'
            payload=torch.load(path,map_location='cpu',weights_only=False)
            changed=[k for k,v in payload['model'].items() if not torch.equal(v,original[k])]
            backbone=[k for k in changed if k.startswith('legacy.backbone.')]
            unexpected=[k for k in backbone if not ('.enc.enc4.' in k or '.dec.' in k)]
            buffers=[k for k in changed if k.endswith(('running_mean','running_var','num_batches_tracked'))]
            if unexpected or buffers:raise AssertionError((unexpected,buffers))
            rows.append(dict(run=run,checkpoint=kind,epoch=payload['epoch'],sha256=digest(path),
                changed_tensors=len(changed),changed_backbone_tensors=len(backbone),
                unexpected_backbone_changes=len(unexpected),changed_normalization_buffers=len(buffers)))
    dump(root/'checkpoint_integrity.json',rows)
    return rows


def benchmark(root,selection):
    output=root/'speed_benchmark.json'
    if output.exists():return json.loads(output.read_text())
    rows=read_manifest(REAL,'val')
    rows=[next(r for r in rows if r['collection']==c) for c in ('CULS','FGI_EMIT','WILDFOREST3D')]
    clouds=[load_npz(r['output']) for r in rows]
    hybrid_selection=json.loads((root/'runs/hybrid_queries/selected.json').read_text())
    models={name:build(path)[0].eval() for name,path in [('baseline',INITIAL),('selected',selection['checkpoint']),
            ('hybrid_queries',root/'runs/hybrid_queries/weights/best.pt')]}
    configs={'baseline':configurations()['legacy_consensus'],'selected':selection['config'],
             'hybrid_queries':hybrid_selection['config']}
    for model in models.values():predict(model,clouds[0])
    measurements=[]
    names=list(models)
    for repeat in range(3):
        for name in names[repeat:]+names[:repeat]:
            for row,a in zip(rows,clouds):
                torch.cuda.synchronize();start=time.monotonic();raw=predict(models[name],a);torch.cuda.synchronize()
                network=time.monotonic()-start;start=time.monotonic()
                labels,_,instances,_=merge_masks_with_sources(a,raw,configs[name]);merge=time.monotonic()-start
                measurements.append(dict(model=name,repeat=repeat,plot=row['dataset_id'],voxels=len(labels),
                    network_seconds=network,merge_seconds=merge,total_seconds=network+merge,instances=len(instances)))
                print(f'speed {name} {repeat+1}/3 {row["collection"]}: {network+merge:.2f}s',flush=True)
    summary={}
    for name in models:
        values=[sum(r['total_seconds'] for r in measurements if r['model']==name and r['repeat']==i) for i in range(3)]
        summary[name]=dict(mean_seconds=float(np.mean(values)),sd_seconds=float(np.std(values,ddof=1)),repeat_seconds=values)
    summary['selected_to_baseline_ratio']=summary['selected']['mean_seconds']/summary['baseline']['mean_seconds']
    summary['hybrid_to_baseline_ratio']=summary['hybrid_queries']['mean_seconds']/summary['baseline']['mean_seconds']
    report=dict(summary=summary,measurements=measurements,device=torch.cuda.get_device_name(),
                scope='Three identical native validation plots, three rotating-order repetitions; warmed network + merging; excludes loading/export')
    dump(output,report)
    del models;torch.cuda.empty_cache()
    return report


def weak_reference_audit(root,selection):
    output=root/'ecodse_weak_2d.json'
    if output.exists():return json.loads(output.read_text())
    rows=[r for r in read_manifest(REAL,'val') if r['collection']=='ECODSE']
    records=[]
    for name,path,corrected in [('baseline_old_heights',INITIAL,False),('baseline_smrf_heights',INITIAL,True),
                                ('selected_smrf_heights',selection['checkpoint'],True)]:
        model,_=build(path);model.eval()
        config=selection['config'] if name.startswith('selected') else configurations()['legacy_consensus']
        for row in rows:
            arrays=load_npz(row['output'] if corrected else row['original_prepared'])
            raw=predict(model,arrays);_,_,instances,_=merge_masks_with_sources(arrays,raw,config)
            gt=gpd.read_file(row['gt_vector'])
            result=metrics(gt.geometry,[p['geometry'] for p in instances])
            records.append(dict(model=name,plot=row['dataset_id'],reference_crowns=len(gt),predictions=len(instances),
                matched_reference_crowns=result['tp'],reference_recall=result['tp']/max(len(gt),1),
                matched_mean_iou=result['iou_sum']/max(result['tp'],1)))
            print(f'weak 2D {name} {row["dataset_id"]}',flush=True)
        del model;torch.cuda.empty_cache()
    report=dict(records=records,used_for_selection=False,
        limitation='Partial polygon labels: reference recall only. No reliable all-scene precision, F1, PQ or 3D truth.')
    dump(output,report);return report


def main():
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--root',type=Path,default=DEFAULT_OUTPUT)
    p.add_argument('--skip-gpu',action='store_true');args=p.parse_args();root=args.root
    torch.set_num_threads(4)
    report=json.loads((root/'final_report.json').read_text());selection=json.loads((root/'selected.json').read_text())
    before=json.loads((root/'final_test/baseline/real/metrics.json').read_text())['fixed']
    after=json.loads((root/'final_test/selected/real/metrics.json').read_text())['fixed']
    ablation,hybrid=architecture_ablation(root,cache_only=args.skip_gpu)
    integrity=checkpoint_audit(root)
    intervals={name:{k:bootstrap(before,result,k) for k in ('point_metrics','metrics')}
               for name,result in [('selected',after),('hybrid_queries',hybrid)]}
    dump(root/'paired_uncertainty.json',intervals)
    speed=benchmark(root,selection) if not args.skip_gpu else (json.loads((root/'speed_benchmark.json').read_text()) if (root/'speed_benchmark.json').exists() else {})
    weak=weak_reference_audit(root,selection) if not args.skip_gpu else (json.loads((root/'ecodse_weak_2d.json').read_text()) if (root/'ecodse_weak_2d.json').exists() else {})
    annotation=[]
    for split in ('train','val','test'):
        for row in read_manifest(REAL,split):
            audit=json.loads(Path(row['output']).with_name('annotation_audit.json').read_text())
            annotation.append(dict(dataset_id=row['dataset_id'],split=split,collection=row['collection'],**audit))
    write_csv(root/'annotation_audit.csv',annotation)
    source=[]
    for name,result in [('baseline',before),('selected',after),('hybrid_queries',hybrid)]:
        for branch in ('point_metrics','metrics'):
            for collection,m in result[branch]['by_source'].items():source.append(dict(model=name,branch=branch,collection=collection,**m))
    write_csv(root/'test_by_collection.csv',source)
    excel(root);book=load_workbook(root/'experiments.xlsx')
    validations=[]
    for run in sorted((root/'runs').iterdir()):
        for path in sorted((run/'validation').glob('epoch_*/metrics.json')):
            for fusion,result in json.loads(path.read_text()).items():
                validations.append(dict(run=run.name,epoch=int(path.parent.name.split('_')[1]),fusion=fusion,
                    point_SB_PQ=result['point_metrics']['source_balanced_pq'],crown_SB_PQ=result['metrics']['source_balanced_pq'],
                    split_points=sum(r['revised_split_points'] for r in result['per_plot']),
                    merge_points=sum(r['revised_merge_points'] for r in result['per_plot'])))
    write_csv(root/'validation_ablations.csv',validations)
    for name,rows in [('annotation_audit',annotation),('checkpoint_integrity',integrity),('test_by_collection',source),('ablation_tests',[ablation]),
                      ('validation_ablations',validations),('timing',speed.get('measurements',[])),('weak_ECODSE',weak.get('records',[]))]:
        sheet=book.create_sheet(name);fields=list(dict.fromkeys(k for r in rows for k in r))
        if not fields:continue
        sheet.append(fields)
        for row in rows:sheet.append([row.get(k) for k in fields])
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
    book.save(root/'experiments.xlsx')
    lines=['# DualCrown3D supervision v4 results','',f'Selected run: `{report["winner"]}`, epoch {selection["epoch"]}, fusion `{selection["fusion"]}`.',
        f'Checkpoint: `{selection["checkpoint"]}`.',f'Validation non-regression acceptance: {selection["acceptance_passed"]}.','',
        'The baseline and selected checkpoint use the SAME corrected annotation protocol. These numbers are not directly comparable to v3 reports.',
        '', '| Model | Domain | Point SB-PQ | Crown SB-PQ | Point F1 | Crown F1 |', '|---|---|---:|---:|---:|---:|']
    for r in report['comparison']+[ablation]:lines.append(f'| {r["model"]} | {r["domain"]} | {r["point_SB_PQ"]:.4f} | {r["crown_SB_PQ"]:.4f} | {r["point_F1"]:.4f} | {r["crown_F1"]:.4f} |')
    if report['winner']=='data_only' and selection['epoch']==0:
        lines+=['','**No improved trained checkpoint was selected.** The winning epoch-0 model retains the original output_20 weights.',
                'Validation acceptance means non-regression, not evidence of improvement. No production exports were replaced.',
                'The hybrid decoder is reported as an experimental ablation, not a promoted replacement.']
    lines+=['','Training used the GPU. 79 real native training plots + 43 synthetic parent plots (129 flight variants).',
        'Primary real evaluation: 14 validation / 24 test plots. The 30 ECODSE plots remain weak 2D references and are excluded from 3D supervision.',
        'The test was historically exposed; one training seed per variant. Bootstrap intervals are conditional plot uncertainty, not multi-seed uncertainty.','',
        '## Paired test change (collection-balanced PQ)','']
    for name,values in intervals.items():
        for k,r in values.items():lines.append(f'- {name}, {k}: {r["delta"]:+.4f}, paired 95% interval [{r["lower_95"]:+.4f}, {r["upper_95"]:+.4f}].')
    if speed:lines+=['',f'Selected/baseline mean network+merging time ratio: {speed["summary"]["selected_to_baseline_ratio"]:.3f}.']
    if speed:lines+=[f'Experimental hybrid/baseline time ratio: {speed["summary"]["hybrid_to_baseline_ratio"]:.3f}.']
    lines+=['','See `experiments.xlsx` for configuration/checkpoint logs, epochs, per-collection test results, annotation audit and timing.',
            'See `ecodse_weak_2d.json` for supplemental reference-crown recall; it is not an all-scene accuracy estimate.']
    (root/'RESULTS.md').write_text('\n'.join(lines)+'\n')
    print('\n'.join(lines),flush=True)


if __name__=='__main__':main()
