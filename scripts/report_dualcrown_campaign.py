#!/usr/bin/env python3
"""Create paired uncertainty estimates and figures from completed evaluations."""
import argparse
from collections import defaultdict
import json
from pathlib import Path
import sys

import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from openpyxl import load_workbook

PROJECT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(PROJECT))
from scripts.train_dualcrown_campaign import dump, workbook


def paired_interval(before, after, branch, samples=2000):
    left={r['dataset_id']:r for r in before['per_plot']}
    right={r['dataset_id']:r for r in after['per_plot']}
    if left.keys()!=right.keys(): raise ValueError('Paired plots differ')
    groups=defaultdict(list)
    for key,r in left.items(): groups[(r['source_dataset'],r['collection'])].append(key)
    rng=np.random.default_rng(1029)
    def score(rows):
        vals=[r['point_metrics'] if branch=='point' else r for r in rows]
        tp=sum(v['tp'] for v in vals); fp=sum(v['fp'] for v in vals); fn=sum(v['fn'] for v in vals)
        return sum(v['iou_sum'] for v in vals)/max(tp+.5*(fp+fn),1)
    observed=np.mean([score([right[k] for k in keys])-score([left[k] for k in keys]) for keys in groups.values()])
    diffs=[]
    for _ in range(samples):
        values=[]
        for keys in groups.values():
            draw=rng.choice(keys,len(keys),replace=True)
            values.append(score([right[k] for k in draw])-score([left[k] for k in draw]))
        diffs.append(float(np.mean(values)))
    lo,hi=np.quantile(diffs,[.025,.975])
    return dict(delta=float(observed),lower_95=float(lo),upper_95=float(hi),bootstrap_samples=samples,
                protocol='paired plot bootstrap stratified by source collection; source-balanced PQ',
                caveat='Conditional on the selected run; small/single-plot strata limit uncertainty estimation')


def main():
    parser=argparse.ArgumentParser(); parser.add_argument('--root',type=Path,default=PROJECT/'outputs/dualcrown3d_joint_campaign_v1')
    args=parser.parse_args(); root=args.root
    report=json.loads((root/'final_report.json').read_text())
    winner=report['representative']['name']
    uncertainty={}
    for domain in ('real','helios'):
        before=json.loads((root/f'final_test/baseline/{domain}/test_metrics.json').read_text())
        after=json.loads((root/f'final_test/{winner}/{domain}/test_metrics.json').read_text())
        uncertainty[domain]={b:paired_interval(before,after,b) for b in ('point','crown')}
    dump(root/'paired_uncertainty.json',uncertainty)
    trials=[]
    for path in sorted(root.glob('runs/*/selected.json')):
        d=json.loads(path.read_text())
        trials.append((path.parent.name,d['validation']['point_metrics']['source_balanced_pq'],d['validation']['metrics']['source_balanced_pq']))
    baseline=json.loads((root/'baseline/real_val/val_metrics.json').read_text())
    trials.insert(0,('baseline',baseline['point_metrics']['source_balanced_pq'],baseline['metrics']['source_balanced_pq']))
    fig,ax=plt.subplots(figsize=(11,5)); x=np.arange(len(trials))
    ax.bar(x-.2,[r[1] for r in trials],.4,label='Point PQ')
    ax.bar(x+.2,[r[2] for r in trials],.4,label='Crown PQ')
    ax.set_xticks(x,[r[0] for r in trials],rotation=25,ha='right')
    ax.set_ylabel('Source-balanced PQ @ IoU 0.50'); ax.set_title('Real validation: selected checkpoint per experiment')
    ax.legend(); fig.tight_layout(); fig.savefig(root/'validation_comparison.png',dpi=180); plt.close(fig)
    text=['# DualCrown3D joint campaign results','',f'Selected architecture: `{report["architecture_winner"]["name"]}`.',
          f'Representative checkpoint selected on validation: `{winner}`.',
          '', '| Run | Test domain | Point SB-PQ | Crown SB-PQ | Point pooled F1 | Crown pooled F1 |',
          '|---|---|---:|---:|---:|---:|']
    for r in report['test_rows']:
        text.append(f'| {r["stage"]} | {r["domain"]} | {r["point_source_balanced_PQ"]:.4f} | {r["crown_source_balanced_PQ"]:.4f} | {r["point_F1"]:.4f} | {r["crown_F1"]:.4f} |')
    text+=['','All five requested steps were executed. The real test is historically exposed; interpret it as a paired regression benchmark.',
           'All variants of the same HELIOS parent plot remain training-only. Full crown polygons were copied without alteration.',
           'Baseline and candidate evaluations use fixed serialization at every pooling stage. Earlier stochastic evaluations are not directly comparable.',
           'See `paired_uncertainty.json` for stratified paired bootstrap intervals and `final_report.json` for three-seed mean/standard deviation.','']
    configuration=json.loads((root/'runs'/winner/'configuration.json').read_text())
    text += ['## Selected model','',
             f'Checkpoint: `{report["representative"]["checkpoint"]}`.',
             f'Best epoch: {report["representative"]["best_epoch"]}; scope: `{configuration["scope"]}`.',
             f'Decoder: `{json.dumps(configuration["model_args"],sort_keys=True)}`.',
             f'Real-data sampling fraction: {configuration["real_fraction"]:.0%}.','',
             '## Three-seed stability','',
             '| Test domain | Metric | Mean | Sample standard deviation |',
             '|---|---|---:|---:|']
    for domain,metrics in report['three_seed_statistics'].items():
        for metric,values in metrics.items():
            text.append(f'| {domain} | {metric} | {values["mean"]:.4f} | {values["std"]:.4f} |')
    speed_path=root/'speed_benchmark.json'
    if speed_path.exists():
        speed=json.loads(speed_path.read_text())
        text += ['', '## Speed check','',speed['protocol']+'.',
                 f'Selected/baseline median time ratio: {speed["summary"]["selected_to_baseline_ratio"]:.3f} (lower is faster).']
    (root/'RESULTS.md').write_text('\n'.join(text))
    workbook(root)
    book=load_workbook(root/'experiments.xlsx')
    sheet=book.create_sheet('test_by_collection')
    sheet.append(['run','domain','branch','collection','precision','recall','F1','PQ','SQ','TP','FP','FN'])
    for path in sorted(root.glob('final_test/*/*/test_metrics.json')):
        data=json.loads(path.read_text())
        for branch in ('metrics','point_metrics'):
            for collection,m in data[branch]['by_source'].items():
                sheet.append([path.parents[1].name,path.parent.name,branch,collection,
                              *[m[k] for k in ('precision','recall','f1','pq','sq','tp','fp','fn')]])
    sheet.freeze_panes='A2'; sheet.auto_filter.ref=sheet.dimensions
    sheet=book.create_sheet('paired_uncertainty')
    sheet.append(['domain','branch','delta_SB_PQ','lower_95','upper_95','bootstrap_samples'])
    for domain,branches in uncertainty.items():
        for branch,m in branches.items():
            sheet.append([domain,branch,m['delta'],m['lower_95'],m['upper_95'],m['bootstrap_samples']])
    sheet.freeze_panes='A2'; sheet.auto_filter.ref=sheet.dimensions
    sheet=book.create_sheet('three_seed_stability')
    sheet.append(['domain','metric','mean','sample_standard_deviation','number_of_seeds'])
    for domain,metrics in report['three_seed_statistics'].items():
        for metric,values in metrics.items():
            sheet.append([domain,metric,values['mean'],values['std'],len(report['seeds'])])
    sheet.freeze_panes='A2'; sheet.auto_filter.ref=sheet.dimensions
    book.save(root/'experiments.xlsx')
    print(json.dumps(uncertainty,indent=2))


if __name__=='__main__': main()
