#!/usr/bin/env python3
"""Train and validate the dense point branch, keeping the legacy branch fixed."""
import argparse
import csv
import hashlib
import json
import math
import os
import sys
import time
from collections import Counter
from pathlib import Path

import torch
from torch.utils.data import DataLoader, WeightedRandomSampler
from openpyxl import Workbook

PROJECT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT))
from pointcloud.data import PointCloudCropDataset, move_to_device
from pointcloud.dual_head import DualHeadLitePT, dense_losses


def arguments():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('--manifest', type=Path, default=PROJECT.parent / 'combined_als_crowns_no_rectangles_v2/manifest.csv')
    p.add_argument('--initial-weights', type=Path, default=PROJECT / 'outputs/combined_full_crowns_no_rectangles_v2/selected/best.pt')
    p.add_argument('--cluster-config', type=Path, default=PROJECT / 'outputs/combined_full_crowns_no_rectangles_v2/selected/frozen_selection.json')
    p.add_argument('--output-dir', type=Path, default=PROJECT / 'outputs/dual_head_litept_v3')
    p.add_argument('--epochs', type=int, default=32)
    p.add_argument('--samples-per-epoch', type=int, default=120)
    p.add_argument('--max-points', type=int, default=40000)
    p.add_argument('--val-every', type=int, default=4)
    p.add_argument('--lr', type=float, default=.0003)
    p.add_argument('--workers', type=int, default=2)
    p.add_argument('--device', default='cuda:0')
    p.add_argument('--seed', type=int, default=20260925)
    return p.parse_args()


def epoch_pass(model, loader, device, optimizer=None):
    model.train(optimizer is not None)
    totals, n = Counter(), 0
    tp = fp = fn = 0
    start = time.monotonic()
    with torch.set_grad_enabled(optimizer is not None):
        for batch in loader:
            batch = move_to_device(batch, device)
            inputs = {k: batch[k] for k in ('coord', 'grid_coord', 'feat', 'offset')}
            if optimizer is not None:
                inputs['tree_id'] = batch['tree_id']
            result = model(inputs)
            loss = dense_losses(result, batch)
            if not torch.isfinite(loss['loss']):
                raise FloatingPointError('Non-finite dense-head loss')
            if optimizer is not None:
                optimizer.zero_grad(set_to_none=True)
                loss['loss'].backward()
                torch.nn.utils.clip_grad_norm_(model.point_decoder.parameters(), 5., error_if_nonfinite=True)
                optimizer.step()
            totals.update({k: float(v.detach()) for k, v in loss.items()})
            pred = result['point_semantic_logits'].argmax(1) == 1
            true = batch['tree_id'] > 0
            tp += int((pred & true).sum())
            fp += int((pred & ~true).sum())
            fn += int((~pred & true).sum())
            n += 1
    return {**{k: v / max(n, 1) for k, v in totals.items()},
            'foreground_iou': tp / max(tp + fp + fn, 1),
            'foreground_f1': 2 * tp / max(2 * tp + fp + fn, 1),
            'seconds': time.monotonic() - start}


def write_logs(output, rows, config, reference, selected):
    fields = list(dict.fromkeys(k for row in rows for k in row))
    if rows:
        with (output / 'training_log.csv').open('w', newline='') as f:
            writer = csv.DictWriter(f, fieldnames=fields)
            writer.writeheader()
            writer.writerows(rows)
    wb = Workbook()
    sheet = wb.active
    sheet.title = 'Epoki'
    sheet.append(fields or ['epoch'])
    for row in rows:
        sheet.append([row.get(k) for k in fields])
    sheet = wb.create_sheet('Architektura i parametry')
    for k, v in config.items():
        sheet.append([k, json.dumps(v, default=str) if isinstance(v, (list, dict)) else str(v)])
    sheet = wb.create_sheet('Porownanie walidacja')
    sheet.append(['branch', 'source_balanced_PQ', 'pooled_F1', 'pooled_PQ', 'checkpoint'])
    for name, result, ckpt in [('legacy', reference, config['initial_weights']), ('dense_points', selected, str(output / 'weights/best.pt'))]:
        if result:
            m = result['metrics']
            sheet.append([name, m['source_balanced_pq'], m['f1'], m['pq'], ckpt])
    for sheet in wb:
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
    temp = output / 'experiments.tmp.xlsx'
    wb.save(temp)
    os.replace(temp, output / 'experiments.xlsx')


def main():
    args = arguments()
    torch.manual_seed(args.seed)
    torch.set_num_threads(4)
    torch.set_float32_matmul_precision('high')
    if args.device.startswith('cuda') and not torch.cuda.is_available():
        raise RuntimeError('CUDA explicitly requested but unavailable')
    device = torch.device(args.device)
    output = args.output_dir.resolve()
    output.mkdir(parents=True, exist_ok=True)
    weights = output / 'weights'
    weights.mkdir(exist_ok=True)
    if (weights / 'last.pt').exists():
        raise FileExistsError('Use a fresh experiment directory; existing checkpoint protected')
    train = PointCloudCropDataset(args.manifest, 'train', max_points=args.max_points, repeats=2,
                                  augment=True, seed=args.seed, density_keep_fractions=(.35, .5, .75, 1.))
    val = PointCloudCropDataset(args.manifest, 'val', max_points=args.max_points, repeats=2,
                                augment=False, seed=args.seed + 1)
    if any(r['collection'] == 'IDTREES' for r in train.rows + val.rows):
        raise ValueError('Excluded box-annotation dataset present')
    group_counts = Counter((r['source_dataset'], r['collection']) for r in train.rows)
    sample_weights = [1. / group_counts[(r['source_dataset'], r['collection'])] for r in train.rows] * 2
    sampler = WeightedRandomSampler(sample_weights, args.samples_per_epoch, replacement=True)
    options = dict(batch_size=None, num_workers=args.workers, pin_memory=device.type == 'cuda')
    train_loader = DataLoader(train, sampler=sampler, **options)
    val_loader = DataLoader(val, **options)
    model = DualHeadLitePT().to(device)
    model.initialize_legacy(args.initial_weights)
    optimizer = torch.optim.AdamW(model.point_decoder.parameters(), lr=args.lr, weight_decay=.02)
    scheduler = torch.optim.lr_scheduler.CosineAnnealingLR(optimizer, T_max=args.epochs, eta_min=args.lr * .05)
    payload = json.loads(args.cluster_config.read_text())
    cluster = payload.get('config') or payload['model']['config']
    config = {**{k: str(v) if isinstance(v, Path) else v for k, v in vars(args).items()},
              'architecture': 'Frozen LitePT-S + legacy crown branch + SATv2-inspired tree mask branch (ISA, 3 masked cross-attention layers, asymmetric memory, seed one-to-many supervision)',
              'decoder_reference': 'https://arxiv.org/abs/2606.08206',
              'queries': 96, 'decoder_layers': 3, 'attention_heads': 4, 'memory_tokens': 1024,
              'isa_embedding_dim': 5, 'spatial_seed_warmup_epochs': 3,
              'context_scales_m': [.75, 2., 6.], 'hidden_dim': 128, 'residual_blocks': 2,
              'trainable_parameters': sum(p.numel() for p in model.parameters() if p.requires_grad),
              'total_parameters': sum(p.numel() for p in model.parameters()),
              'initial_sha256': hashlib.sha256(args.initial_weights.read_bytes()).hexdigest(),
              'manifest_sha256': hashlib.sha256(args.manifest.read_bytes()).hexdigest(),
              'legacy_cluster_config': cluster, 'selection': 'validation source-balanced point-instance PQ at IoU .5',
              'test_used': False, 'annotation_caveat': 'Mixed native and polygon-projected point labels; unlabeled points include incomplete annotation.'}
    (output / 'configuration.json').write_text(json.dumps(config, indent=2) + '\n')
    os.environ['SEGMENTATION_EXPERIMENT_ROOT'] = str(output)
    from scripts.evaluate_combined_full_crowns import evaluate_offset_model
    from scripts.evaluate_dual_head import evaluate_masks
    reference = evaluate_offset_model(model, args.manifest, output / 'legacy_validation', config=cluster)
    model.backbone.shuffle_orders = False
    rows, selected, best = [], None, -math.inf
    for epoch in range(1, args.epochs + 1):
        train.set_epoch(epoch)
        model.point_decoder.epoch = epoch
        tr = epoch_pass(model, train_loader, device, optimizer)
        va = epoch_pass(model, val_loader, device)
        row = {'epoch': epoch, **{'train_' + k: v for k, v in tr.items()}, **{'val_' + k: v for k, v in va.items()}}
        row['lr'] = optimizer.param_groups[0]['lr']
        scheduler.step()
        full = None
        if epoch % args.val_every == 0 or epoch == args.epochs:
            full = evaluate_masks(model, args.manifest, output / f'validation/epoch_{epoch:03d}', max_points=args.max_points)
            model.backbone.shuffle_orders = False
            row['val_crown_source_balanced_pq'] = full['metrics']['source_balanced_pq']
            row['val_crown_f1'] = full['metrics']['f1']
            row['val_point_source_balanced_pq'] = full['point_metrics']['source_balanced_pq']
            row['val_point_f1'] = full['point_metrics']['f1']
        checkpoint = dict(format='dual_head_litept_v3', model=model.state_dict(), epoch=epoch,
                          config=config, optimizer=optimizer.state_dict(), scheduler=scheduler.state_dict())
        torch.save(checkpoint, weights / 'last.pt')
        if full and full['point_metrics']['source_balanced_pq'] > best:
            best = full['point_metrics']['source_balanced_pq']
            selected = full
            torch.save(checkpoint, weights / 'best.pt')
            selection = dict(epoch=epoch, config=full['config'], validation=full,
                             checkpoint_sha256=hashlib.sha256((weights / 'best.pt').read_bytes()).hexdigest(),
                             test_used_for_selection=False)
            (output / 'selected.json').write_text(json.dumps(selection, indent=2) + '\n')
        rows.append(row)
        write_logs(output, rows, config, reference, selected)
        print(json.dumps(row), flush=True)
    # Exact tensor equality proves that both backbone weights and BN buffers stayed fixed.
    original = torch.load(args.initial_weights, map_location='cpu', weights_only=False)['model']
    unchanged = all(torch.equal(v.detach().cpu(), original[k]) for k, v in model.legacy.state_dict().items())
    report = dict(legacy_weights_and_buffers_unchanged=unchanged, completed_epochs=args.epochs,
                  selected_validation=selected['metrics'], legacy_validation=reference['metrics'],
                  selected_point_validation=selected['point_metrics'],
                  improves_validation_crown_pq=selected['metrics']['source_balanced_pq'] > reference['metrics']['source_balanced_pq'], test_evaluated=False)
    (output / 'training_report.json').write_text(json.dumps(report, indent=2) + '\n')
    if not unchanged:
        raise AssertionError('Frozen legacy branch unexpectedly changed')
    print(json.dumps(report, indent=2), flush=True)


if __name__ == '__main__':
    main()
