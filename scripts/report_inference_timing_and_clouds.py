#!/usr/bin/env python3
"""Create a compact Excel report from measured inference times and LAZ headers."""

import argparse
import json
import os
from pathlib import Path
import statistics

import laspy
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


PROJECT = Path(__file__).resolve().parents[1]


def add_sheet(book, name, headings, rows):
    sheet = book.create_sheet(name)
    sheet.append(headings)
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="24536C")
        cell.alignment = Alignment(wrap_text=True)
    for column in sheet.columns:
        letter = get_column_letter(column[0].column)
        length = max(len(str(cell.value or "")) for cell in column)
        sheet.column_dimensions[letter].width = min(max(length + 2, 13), 65)
    return sheet


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--output-dir", type=Path,
        default=PROJECT / "output_20_dualcrown3d_joint_finetune",
    )
    parser.add_argument("--overwrite", action="store_true", help="Replace this report only")
    args = parser.parse_args()
    output_dir = args.output_dir.resolve()
    timing_path = output_dir / "inference_timing_5runs.json"
    timing = json.loads(timing_path.read_text())
    inference = json.loads((output_dir / "inference_report.json").read_text())
    preparation_path = Path(timing["prepared_dir"]) / "preparation.json"
    preparation = json.loads(preparation_path.read_text())

    tiles = {tile["tile_id"]: tile for tile in preparation["tiles"]}
    clouds = {}
    for record in inference["clouds"]:
        path = Path(record["file"])
        tile_id = path.stem.removeprefix("trees_")
        if tile_id in clouds or tile_id not in tiles:
            raise ValueError(f"Unexpected/duplicate cloud tile: {tile_id}")
        clouds[tile_id] = (path, record)
    if set(clouds) != set(tiles):
        raise ValueError("Output clouds do not match prepared core tiles")

    measurements = timing["measurements"]
    durations = [row["wall_seconds"] for row in measurements]
    area_ha = sum(
        (tile["bounds"][2] - tile["bounds"][0])
        * (tile["bounds"][3] - tile["bounds"][1]) / 10_000
        for tile in tiles.values()
    )
    if len(measurements) != timing["runs"] or len(measurements) < 2:
        raise ValueError("Incomplete timing measurements")
    if abs(area_ha - timing["area_ha"]) > 1e-8:
        raise ValueError("Timing area differs from the prepared core tile area")
    if abs(statistics.mean(durations) - timing["summary"]["mean_seconds"]) > 1e-6:
        raise ValueError("Saved timing summary does not match measurements")
    if timing["checkpoint_sha256"] != inference["checkpoint_sha256"]:
        raise ValueError("Timing and output cloud checkpoints differ")

    point_count = sum(record["points"] for _, record in clouds.values())
    labelled_count = sum(record["labelled_points"] for _, record in clouds.values())
    mean_seconds = statistics.mean(durations)
    mean_seconds_per_ha = mean_seconds / area_ha
    book = Workbook()
    book.remove(book.active)
    summary_rows = [
        ("Liczba powtórzeń", len(measurements), "", "Zmierzona inferencja całego obszaru"),
        ("Obszar czterech kafli", area_ha, "ha", "Suma granic kafli CHM z preparation.json"),
        ("Liczba punktów w wynikowych LAZ", point_count, "pkt", "Suma raportów eksportu"),
        ("Liczba punktów źródłowego ALS", preparation["raw_points_read"], "pkt", "Odczytane ze źródłowego LAZ"),
        ("Punkty ALS w obszarze przygotowania", preparation["points_in_extent"], "pkt", "Obszar z buforem, przed przypisaniem do kafli"),
        ("Punkty gruntu w ALS", preparation["ground_points"], "pkt", "Klasa 2 użyta do normalizacji wysokości"),
        ("Liczba punktów z tree_id", labelled_count, "pkt", "Suma raportów eksportu"),
        ("Udział punktów z tree_id", labelled_count / point_count, "ułamek", "Dotyczy wszystkich punktów, także gruntu"),
        ("Średnia gęstość punktów", point_count / (area_ha * 10_000), "pkt/m²", "Powierzchnia kafli, nie obwiednia chmury"),
        ("Liczba wokseli wejściowych", timing["voxel_count"], "woksele", "Wszystkie punkty przetwarzane razem"),
        ("Średni czas", mean_seconds, "s", "Średnia arytmetyczna pięciu pomiarów wall-clock"),
        ("Odchylenie standardowe czasu", statistics.stdev(durations), "s", "Odchylenie próbki (n-1)"),
        ("Mediana czasu", statistics.median(durations), "s", ""),
        ("Najkrótszy czas", min(durations), "s", ""),
        ("Najdłuższy czas", max(durations), "s", ""),
        ("Średni czas na hektar", mean_seconds_per_ha, "s/ha", "Cały obszar; nie osobny pomiar każdego kafla"),
        ("Odchylenie standardowe na hektar", statistics.stdev(durations) / area_ha, "s/ha", "Odchylenie próbki (n-1)"),
        ("Czas pojedynczego wcześniejszego uruchomienia", timing["original_inference_seconds"], "s", "Nie należy do pięciu powtórzeń"),
        ("Urządzenie", timing["device"], "", "CUDA GPU"),
        ("Checkpoint SHA256", timing["checkpoint_sha256"], "", ""),
    ]
    summary = add_sheet(book, "Podsumowanie", ("Parametr", "Wartość", "Jednostka", "Uwagi"), summary_rows)
    summary.cell(9, 2).number_format = "0.00%"
    summary.cell(10, 2).number_format = "0.0000"

    measurements_rows = [
        (row["run"], row["wall_seconds"], row["seconds_per_hectare"],
         row["internal_predict_seconds"], row["windows"], row["candidate_masks"])
        for row in measurements
    ]
    measured = add_sheet(
        book, "Pomiary", ("Powtórzenie", "Czas wall-clock [s]", "Czas [s/ha]",
                           "Czas wewnętrzny predict [s]", "Okna", "Kandydaci masek"),
        measurements_rows,
    )
    for row in measured.iter_rows(min_row=2):
        for cell in row[1:4]:
            cell.number_format = "0.000"

    cloud_rows = []
    for tile_id in sorted(clouds):
        path, record = clouds[tile_id]
        bounds = tiles[tile_id]["bounds"]
        tile_area_ha = (bounds[2] - bounds[0]) * (bounds[3] - bounds[1]) / 10_000
        with laspy.open(path) as reader:
            header = reader.header
            if header.point_count != record["points"]:
                raise ValueError(f"Point count mismatch: {path}")
            crs = header.parse_crs()
            if crs is None or crs.to_epsg() != 2180:
                raise ValueError(f"Unexpected cloud CRS: {path}")
            cloud_rows.append((
                tile_id, str(path), tile_area_ha, record["points"],
                record["points"] / (tile_area_ha * 10_000),
                record["labelled_points"], record["labelled_points"] / record["points"],
                record["trees"], path.stat().st_size / 1_000_000,
                header.mins[0], header.mins[1], header.mins[2],
                header.maxs[0], header.maxs[1], header.maxs[2],
                f"EPSG:{crs.to_epsg()}", str(header.version), header.point_format.id,
                ", ".join(header.point_format.extra_dimension_names),
            ))
    cloud_sheet = add_sheet(
        book, "Chmury", ("Kafel", "Plik LAZ", "Obszar kafla [ha]", "Punkty", "Gęstość [pkt/m²]",
                           "Punkty z tree_id", "Udział z tree_id", "Liczba tree_id", "Plik [MB]",
                           "X min [m]", "Y min [m]", "Z min [m]",
                           "X max [m]", "Y max [m]", "Z max [m]",
                           "CRS", "LAS", "Format punktu", "Dodatkowe pola"),
        cloud_rows,
    )
    for row in cloud_sheet.iter_rows(min_row=2):
        row[6].number_format = "0.00%"
        for cell in (row[2], row[4], row[8]):
            cell.number_format = "0.000"
        for cell in row[9:15]:
            cell.number_format = "0.00"

    source_rows = []
    for source in preparation["als_files"]:
        source_path = Path(source)
        with laspy.open(source_path) as reader:
            header = reader.header
            try:
                source_crs = header.parse_crs()
                source_crs_text = str(source_crs) if source_crs is not None else "Brak w nagłówku"
            except Exception as error:
                source_crs_text = f"Nieprawidłowe WKT ({type(error).__name__})"
            source_rows.append((
                str(source_path), source_path.stat().st_size / 1_000_000,
                header.point_count, header.mins[0], header.mins[1], header.mins[2],
                header.maxs[0], header.maxs[1], header.maxs[2],
                source_crs_text, preparation["chm_crs"], str(header.version),
                header.point_format.id,
            ))
    add_sheet(
        book, "Zrodlo_ALS", ("Plik", "Rozmiar [MB]", "Punkty", "X min [m]", "Y min [m]",
                             "Z min [m]", "X max [m]", "Y max [m]", "Z max [m]",
                             "CRS z nagłówka", "CRS użyty wg CHM", "LAS", "Format punktu"),
        source_rows,
    )

    methodology_rows = [
        ("Źródło czasów", str(timing_path)),
        ("Źródło metadanych chmur", str(output_dir / "inference_report.json")),
        ("Źródło obszarów", str(preparation_path)),
        ("Źródłowa chmura ALS", ", ".join(preparation["als_files"])),
        ("CRS źródła", "Źródłowy LAZ ma niepoprawne/puste WKT; EPSG:2180 pochodzi z kafli CHM i jest zapisany w wynikowych LAZ."),
        ("Konfiguracja", f"Woksel {inference['voxel_size']} m; okno {inference['window_size']} m; nakładanie {inference['overlap']} m"),
        ("Zakres pomiaru", timing["protocol"]),
        ("Gęstość", "Liczba punktów wynikowego LAZ / powierzchnia granic kafla CHM; nie rzeczywista powierzchnia koron."),
        ("Z min/max", "Bezwzględne wysokości z nagłówka LAZ, nie height_agl."),
        ("Czas/kafel", "Nie był mierzony osobno; czasy z arkusza Pomiary dotyczą czterech kafli łącznie."),
        ("Eksport", "Czas pomiaru nie obejmuje scalania masek, zapisu LAZ ani poligonów."),
    ]
    add_sheet(book, "Metodyka", ("Pole", "Opis"), methodology_rows)
    destination = output_dir / "raport_czasu_inferencji_i_chmur_punktowych.xlsx"
    if destination.exists() and not args.overwrite:
        raise FileExistsError(f"Refusing to overwrite existing report: {destination}")
    temporary = destination.with_suffix(".tmp.xlsx")
    book.save(temporary)
    os.replace(temporary, destination)
    print(destination)


if __name__ == "__main__":
    main()
