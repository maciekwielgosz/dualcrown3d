"""Atomic JSON records and a locally readable Excel experiment workbook."""
import csv
import fcntl
import json
import os
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

ROOT = Path(os.environ.get("SEGMENTATION_EXPERIMENT_ROOT", Path(__file__).resolve().parents[1] / "outputs/combined_full_crowns_v1")).resolve()


def record(experiment_id, **values):
    ROOT.mkdir(parents=True, exist_ok=True)
    records = ROOT / "records"
    records.mkdir(exist_ok=True)
    with (ROOT / ".registry.lock").open("a") as lock:
        fcntl.flock(lock, fcntl.LOCK_EX)
        path = records / f"{experiment_id}.json"
        payload = json.loads(path.read_text()) if path.exists() else {}
        for field in ("val_metrics", "latest_val_metrics", "selected_val_metrics", "final_val_metrics", "test_metrics"):
            if field in values:
                prefix = field.removesuffix("_metrics")
                values.update({f"{prefix}_{key}": value for key, value in values[field].items() if not isinstance(value, (list, dict))})
        payload.update(experiment_id=experiment_id, updated_utc=datetime.now(timezone.utc).isoformat(), **values)
        tmp = path.with_suffix(".tmp")
        tmp.write_text(json.dumps(payload, indent=2, default=str, allow_nan=False) + "\n")
        os.replace(tmp, path)
        rebuild()


def add_sheet(workbook, name, rows):
    sheet = workbook.create_sheet(name)
    if not rows:
        return
    fields = list(dict.fromkeys(key for row in rows for key in row))
    if name == "Eksperymenty":
        priority = ["experiment_id", "status", "architecture", "protocol_version", "grid_version", "final_val_source_balanced_pq", "final_val_f1", "val_source_balanced_pq", "test_source_balanced_pq", "test_f1", "test_recall", "target_met_on_test", "final_validation_checkpoint", "checkpoint", "epoch"]
        fields = [k for k in priority if k in fields] + [k for k in fields if k not in priority]
    sheet.append(fields)
    for row in rows:
        values = []
        for key in fields:
            value = row.get(key)
            if isinstance(value, (dict, list)):
                value = json.dumps(value, default=str)
            elif name == "Epoki" and isinstance(value, str) and key != "experiment_id":
                try:
                    value = float(value) if value else None
                except ValueError:
                    pass
            values.append(value)
        sheet.append(values)
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="254E70")
        sheet.column_dimensions[cell.column_letter].width = min(52, max(16, len(str(cell.value))+3))
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top")
            if isinstance(cell.value, str) and cell.value.startswith("/home/") and len(cell.value) < 400:
                cell.hyperlink = Path(cell.value).as_uri()
                cell.font = Font(color="0563C1", underline="single")


def rebuild():
    workbook = Workbook()
    workbook.remove(workbook.active)
    records = [json.loads(p.read_text()) for p in sorted((ROOT / "records").glob("*.json"))]
    per_source = []
    parameters = []
    for item in records:
        selected = Path(item.get("training_dir", "__missing__")) / "selected_validation.json"
        if selected.exists():
            item["selected_val_metrics"] = json.loads(selected.read_text())["metrics"]
        for field in ("val_metrics", "selected_val_metrics", "final_val_metrics", "test_metrics"):
            for key, value in item.get(field, {}).items():
                if isinstance(value, (float, int)):
                    item[field.removesuffix("_metrics")+"_"+key] = value
            for source, values in item.get(field, {}).get("by_source", {}).items():
                per_source.append(dict(experiment_id=item["experiment_id"], split=field, source=source, **values))
        parameters.extend(dict(experiment_id=item["experiment_id"], parameter=k, value=v) for k,v in item.get("parameters", {}).items())
    add_sheet(workbook, "Eksperymenty", records)
    add_sheet(workbook, "Parametry", parameters)
    add_sheet(workbook, "Metryki_zrodla", per_source)
    epochs = []
    for item in records:
        log = Path(item.get("training_dir", "__missing__")) / "training_log.csv"
        if log.exists():
            with log.open(newline="") as stream:
                epochs.extend({"experiment_id": item["experiment_id"], **r} for r in csv.DictReader(stream))
    add_sheet(workbook, "Epoki", epochs)
    postprocessing = []
    for path in ROOT.glob("*/postprocessing_trials.json"):
        for trial in json.loads(path.read_text()):
            postprocessing.append(dict(experiment_id=trial["experiment_id"], trial=trial["trial"],
                                       **trial["config"], **{k:v for k,v in trial["metrics"].items() if not isinstance(v, dict)},
                                       checkpoint_sha256=trial["checkpoint_sha256"], protocol_version=trial["protocol_version"]))
    if postprocessing:
        add_sheet(workbook, "Postprocessing", postprocessing)
    dataset = Path(os.environ.get("SEGMENTATION_DATA_ROOT", ROOT.parents[2] / "combined_als_crowns_v1")).resolve()
    split = dataset / "split_summary.json"
    if split.exists():
        add_sheet(workbook, "Podzial", [{"pole": k, "wartosc": v} for k, v in json.loads(split.read_text()).items()])
    historical = ROOT / "historical_baselines.json"
    if historical.exists():
        add_sheet(workbook, "Historyczne_run_r", json.loads(historical.read_text()))
    filtered_campaign = (dataset / "exclusion_report.json").exists()
    add_sheet(workbook, "Protokol", [
        {"temat": "Cel", "opis": "Source-balanced PQ@IoU0.50 pełnych koron; dodatkowo F1, precision, recall, SQ. Wybor modelu tylko na val."},
        {"temat": "Poligony", "opis": "Oficjalne pełne obrysy dla crown_overlay; udokumentowana projekcja oznaczonych punktów dla point_native. Projekcja nie jest niezależną ręczną prawdą referencyjną."},
        {"temat": "Test", "opis": "Nowy zamrożony podział przestrzenny; historycznie część obszarów była już oglądana. Nie deklarujemy dziewiczego testu publikacyjnego."},
        {"temat": "Porównanie", "opis": "Wyniki historyczne mają inne zbiory/GT i nie są progiem sukcesu nowego eksperymentu. Baseline przeliczany na tym samym nowym zbiorze."},
        {"temat": "Czas", "opis": "Raportuj osobno inferencję punktową i cały eksport. Crop IoU nie jest polygon F1/PQ."},
        {"temat": "Wyniki robocze", "opis": ("Kampania bez IDTREES od początku używa poprawionej siatki i reguły out-points. Porównuj final_val_* oraz test_*; val_* pokazuje przebieg treningu."
                                                    if filtered_campaign else
                                                    "Epoki i selected_val_* zawierają wyniki sprzed korekty siatki/obszarów bez anotacji. Porównuj final_val_* oraz test_*. Zachowano stare wyniki dla audytu, nie jako wynik końcowy.")},
        {"temat": "Wykluczenie", "opis": ("IDTREES usunięto w całości: 437 plików z prostokątnymi anotacjami. Pozostały podział zachowuje grupy i przydziały v1."
                                               if filtered_campaign else "Brak wykluczenia całej kolekcji w tej kampanii.")},
        {"temat": "Obszary bez anotacji", "opis": "Po dopasowaniu GT nie liczymy FP, gdy >=50% powierzchni niepasującej korony leży w projekcji jawnej klasy 3 (out-points), bez innych niegruntowych punktów w komórce 0.5 m. Duplikaty GT nadal są FP. Poligony nie są przycinane; maska tylko do oceny, wspólna dla DL i baseline."},
        {"temat": "Ograniczenia uczenia", "opis": "Klasa 3 była tłem podczas treningu; maska ignore koryguje ocenę, nie naprawia tego ograniczenia nadzoru. Nie zakładamy, że nieoznaczone drzewa w źródłach crown-overlay mają wyczerpujące adnotacje."},
    ])
    temporary = ROOT / "experiments.tmp.xlsx"
    workbook.save(temporary)
    os.replace(temporary, ROOT / "experiments.xlsx")
